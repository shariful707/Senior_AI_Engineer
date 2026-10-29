"""Create synthetic, explicitly labeled fixtures. No private/customer data."""
import io
import json
from pathlib import Path

import pymupdf
import xlwt
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont

ROOT = Path(__file__).parent / 'fixtures'
ROOT.mkdir(exist_ok=True)
cases = []


def invoice(n, **changes):
    result = {'invoice_number': f'INV-{1000+n}', 'invoice_date': '2026-09-14', 'supplier_name': f'Supplier {n} Ltd', 'customer_name': 'Northwind Traders', 'subtotal': 1200.0, 'tax': 144.0, 'total': 1344.0, 'currency': 'USD', 'line_items': [{'description': 'Widget A', 'quantity': 10, 'unit_price': 120, 'line_total': 1200}]}
    result.update(changes)
    return result


def text_for(record):
    labels = {'invoice_number': 'Invoice number', 'invoice_date': 'Invoice date', 'supplier_name': 'Supplier', 'customer_name': 'Customer', 'subtotal': 'Subtotal', 'tax': 'Tax', 'total': 'Total', 'currency': 'Currency'}
    lines = ['INVOICE', ''] + [f'{label}: {record[key]}' for key, label in labels.items() if record[key] is not None]
    lines.extend(['', 'Description | Quantity | Unit price | Line total'])
    lines.extend(f"{i['description']} | {i['quantity']} | {i['unit_price']} | {i['line_total']}" for i in record['line_items'])
    return '\n'.join(lines)


def add(name, kind, record=None, **extra):
    cases.append(dict(filename=name, category=kind, expected=record, **extra))


def pdf(name, record, pages=1, scanned=False):
    doc = pymupdf.open()
    text = text_for(record)
    for n in range(pages):
        page = doc.new_page(width=595, height=842)
        if scanned:
            image = render(text)
            stream = io.BytesIO()
            image.save(stream, format='PNG')
            page.insert_image(page.rect, stream=stream.getvalue())
        else:
            page.insert_text((45, 60), text if n == 0 else 'Supporting document\nDelivery confirmed.\nThank you for your business.', fontsize=12)
    doc.save(ROOT / name, deflate=True)
    doc.close()


def render(text):
    image = Image.new('RGB', (1600, 1800), 'white')
    font_paths = [Path('C:/Windows/Fonts/consola.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf')]
    font = next((ImageFont.truetype(str(p), 28) for p in font_paths if p.exists()), ImageFont.load_default(size=28))
    ImageDraw.Draw(image).multiline_text((70, 80), text, fill='black', font=font, spacing=17)
    return image


for n in range(1, 7):
    rec = invoice(n)
    if n == 3:
        rec = invoice(n, total=1400)
    if n == 4:
        rec = invoice(n, invoice_number=None, invoice_date=None, currency=None)
    name = f'digital-{n}.pdf'
    pdf(name, rec, pages=3 if n == 2 else 1)
    add(name, 'digital_pdf', rec, needs_review=n in (3, 4))

for n, ext in [(7, 'png'), (8, 'jpg'), (9, 'tiff')]:
    rec = invoice(n)
    image = render(text_for(rec))
    image.save(ROOT / f'image-{n}.{ext}')
    add(f'image-{n}.{ext}', 'image_ocr', rec)

rec = invoice(10)
pdf('scanned.pdf', rec, scanned=True)
add('scanned.pdf', 'scanned_pdf', rec)
rec = invoice(11)
render(text_for(rec)).rotate(2, fillcolor='white').filter(ImageFilter.GaussianBlur(1)).save(ROOT / 'poor-quality.png')
add('poor-quality.png', 'poor_quality', rec)

header = ['Invoice number', 'Supplier', 'Invoice date', 'Total', 'Currency']
records = [invoice(1), invoice(2), invoice(6, total=1390)]
expected_ledger = [{k: r[k] for k in ('invoice_number', 'supplier_name', 'invoice_date', 'total', 'currency')} for r in records]
workbook = Workbook()
for name, subset in [('September', records[:2]), ('Exceptions', records[2:])]:
    sheet = workbook.active if name == 'September' else workbook.create_sheet()
    sheet.title = name
    sheet.append(header)
    for r in subset:
        sheet.append([r['invoice_number'], r['supplier_name'], r['invoice_date'], r['total'], r['currency']])
workbook.save(ROOT / 'ledger.xlsx')
add('ledger.xlsx', 'spreadsheet', ledger_rows=expected_ledger)

import csv

with (ROOT / 'ledger.csv').open('w', newline='', encoding='utf-8') as file:
    writer = csv.writer(file)
    writer.writerow(header)
    for r in records:
        writer.writerow([r['invoice_number'], r['supplier_name'], r['invoice_date'], r['total'], r['currency']])
add('ledger.csv', 'spreadsheet', ledger_rows=expected_ledger)

legacy = xlwt.Workbook()
sheet = legacy.add_sheet('Legacy ledger')
for col, value in enumerate(header):
    sheet.write(0, col, value)
for row, r in enumerate(records, 1):
    for col, key in enumerate(['invoice_number', 'supplier_name', 'invoice_date', 'total', 'currency']):
        sheet.write(row, col, r[key])
legacy.save(str(ROOT / 'ledger.xls'))
add('ledger.xls', 'spreadsheet', ledger_rows=expected_ledger)

(ROOT / 'corrupted.pdf').write_bytes(b'%PDF-1.7\nNot a valid document')
add('corrupted.pdf', 'corrupted', expected_failure=True)
blank = pymupdf.open()
blank.new_page()
blank.save(ROOT / 'blank.pdf')
blank.close()
add('blank.pdf', 'blank', expected_failure=True)

rec = invoice(1, total=1500, tax=300)
pdf('conflicting-copy.pdf', rec)
add('conflicting-copy.pdf', 'digital_pdf', rec, needs_review=True)
rec = invoice(12, total=-10, subtotal=-10, tax=0, line_items=[])
pdf('negative.pdf', rec)
add('negative.pdf', 'digital_pdf', rec, needs_review=True)

pdf('large-1000-pages.pdf', invoice(20), pages=1000)
add('large-1000-pages.pdf', 'large_pdf', invoice(20))
(ROOT / 'manifest.json').write_text(json.dumps(cases, indent=2), encoding='utf-8')
print(f'Generated {len(cases)} labeled documents in {ROOT}')
