"""Bounded, page-aware document parsing; no LLM or database dependencies."""
import csv
import io
import os
import re
import shutil
import zipfile
from datetime import date, datetime
from pathlib import Path

import pymupdf
import pytesseract
import xlrd
from openpyxl import load_workbook
from PIL import Image, ImageOps, ImageSequence

MAX_PAGES = int(os.getenv('MAX_PAGES', '1200'))
MAX_ROWS = int(os.getenv('MAX_SHEET_ROWS', '10000'))
MAX_CHARS = int(os.getenv('MAX_TEXT_CHARS', '2000000'))
Image.MAX_IMAGE_PIXELS = 40_000_000
cmd = os.getenv('TESSERACT_CMD') or shutil.which('tesseract')
if not cmd and Path(r'C:\Program Files\Tesseract-OCR\tesseract.exe').exists():
    cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
if cmd:
    pytesseract.pytesseract.tesseract_cmd = cmd


def detect_type(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip('.')
    ext = {'jpeg': 'jpg', 'tif': 'tiff'}.get(ext, ext)
    magic = None
    if data.startswith(b'%PDF-'):
        magic = 'pdf'
    elif data.startswith(b'\x89PNG\r\n\x1a\n'):
        magic = 'png'
    elif data.startswith(b'\xff\xd8\xff'):
        magic = 'jpg'
    elif data[:4] in (b'II*\x00', b'MM\x00*'):
        magic = 'tiff'
    elif data.startswith(b'PK\x03\x04'):
        magic = 'xlsx'
    elif data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        magic = 'xls'
    elif ext == 'csv' and b'\x00' not in data:
        magic = 'csv'
    if not magic or magic != ext:
        raise ValueError('File signature does not match a supported extension.')
    return magic


def ocr(image: Image.Image, page_number: int) -> dict:
    if image.width * image.height > 40_000_000:
        raise ValueError('Image exceeds the 40 megapixel limit.')
    image = ImageOps.autocontrast(ImageOps.grayscale(ImageOps.exif_transpose(image)))
    data = pytesseract.image_to_data(image, config='--psm 6', timeout=45, output_type=pytesseract.Output.DICT)
    words, confidences = [], []
    last_line = None
    for i, word in enumerate(data['text']):
        if not word.strip():
            continue
        line = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        words.append(('\n' if last_line is not None and line != last_line else ' ') + word)
        last_line = line
        if float(data['conf'][i]) >= 0:
            confidences.append(float(data['conf'][i]) / 100)
    return {'page_number': page_number, 'raw_text': ''.join(words).strip(), 'ocr_confidence': sum(confidences) / len(confidences) if confidences else 0.0, 'method': 'ocr', 'sheet_name': None}


def cell_text(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return '' if value is None else str(value).strip()


ALIASES = {
    'invoice_number': {'invoice_number', 'invoice_no', 'invoice', 'invoice_id'},
    'supplier_name': {'supplier_name', 'supplier', 'vendor', 'vendor_name'},
    'invoice_date': {'invoice_date', 'date'},
    'total': {'total', 'amount', 'gross_total', 'invoice_total'},
    'currency': {'currency', 'currency_code'},
}


def sheet_result(name: str, rows: list) -> tuple:
    if len(rows) > MAX_ROWS:
        raise ValueError(f'Sheet {name} exceeds {MAX_ROWS} rows.')
    text = '\n'.join(' | '.join(cell_text(c) for c in row) for row in rows)
    ledger = []
    header_index, mapping = None, {}
    for row_index, row in enumerate(rows[:20]):
        headers = [re.sub(r'[^a-z0-9]+', '_', cell_text(c).lower()).strip('_') for c in row]
        candidate = {field: next((i for i, h in enumerate(headers) if h in names), None) for field, names in ALIASES.items()}
        if candidate['invoice_number'] is not None and candidate['total'] is not None:
            header_index, mapping = row_index, candidate
            break
    if header_index is not None:
        for n, row in enumerate(rows[header_index + 1:], header_index + 2):
            record = {field: cell_text(row[i]) if i is not None and i < len(row) else None for field, i in mapping.items()}
            if not record['invoice_number']:
                continue
            raw = record['total']
            try:
                record['total'] = float(re.sub(r'[^0-9.\-]', '', raw)) if raw else None
            except ValueError:
                record['total'] = None
            record.update(sheet_name=name, row_number=n)
            ledger.append(record)
    return {'page_number': None, 'raw_text': f'Sheet: {name}\n{text}', 'ocr_confidence': None, 'method': 'spreadsheet', 'sheet_name': name}, ledger


def parse_document(data: bytes, filename: str) -> dict:
    kind = detect_type(data, filename)
    pages, ledger = [], []
    if kind == 'pdf':
        with pymupdf.open(stream=data, filetype='pdf') as doc:
            if doc.needs_pass:
                raise ValueError('Password-protected PDF is not supported.')
            if len(doc) > MAX_PAGES:
                raise ValueError(f'PDF exceeds {MAX_PAGES} pages.')
            chars = 0
            for n, page in enumerate(doc):
                text = page.get_text(sort=True).strip()
                if len(re.sub(r'\s', '', text)) >= 20:
                    result = {'page_number': n + 1, 'raw_text': text, 'ocr_confidence': None, 'method': 'pdf_text', 'sheet_name': None}
                else:
                    if page.rect.width * page.rect.height * (150 / 72) ** 2 > 40_000_000:
                        raise ValueError('PDF page exceeds rendering pixel limit.')
                    pix = page.get_pixmap(dpi=150, colorspace=pymupdf.csRGB)
                    result = ocr(Image.frombytes('RGB', (pix.width, pix.height), pix.samples), n + 1)
                chars += len(result['raw_text'])
                if chars > MAX_CHARS:
                    raise ValueError('Document text exceeds configured character limit; split the document.')
                pages.append(result)
    elif kind in ('png', 'jpg', 'tiff'):
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError('Image exceeds the 40 megapixel limit.')
            for n, frame in enumerate(ImageSequence.Iterator(image)):
                if n >= MAX_PAGES:
                    raise ValueError('Image exceeds page limit.')
                pages.append(ocr(frame.copy(), n + 1))
    elif kind == 'csv':
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(text), dialect):
            if len(rows) >= MAX_ROWS:
                raise ValueError('CSV exceeds row limit.')
            rows.append(row)
        page, ledger = sheet_result('CSV', rows)
        pages.append(page)
    elif kind == 'xlsx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Workbook decompressed size exceeds 100 MB.')
            if '[Content_Types].xml' not in archive.namelist() or 'xl/workbook.xml' not in archive.namelist():
                raise ValueError('ZIP file is not an Excel workbook.')
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if len(workbook.worksheets) > MAX_PAGES:
                raise ValueError('Workbook exceeds sheet limit.')
            for sheet in workbook:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= MAX_ROWS:
                        raise ValueError('Workbook exceeds row limit.')
                    rows.append(row)
                page, records = sheet_result(sheet.title, rows)
                pages.append(page)
                ledger.extend(records)
        finally:
            workbook.close()
    elif kind == 'xls':
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            if workbook.nsheets > MAX_PAGES:
                raise ValueError('Workbook exceeds sheet limit.')
            for sheet in workbook.sheets():
                if sheet.nrows > MAX_ROWS:
                    raise ValueError('Workbook exceeds row limit.')
                rows = []
                for n in range(sheet.nrows):
                    rows.append([xlrd.xldate.xldate_as_datetime(c.value, workbook.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value for c in sheet.row(n)])
                page, records = sheet_result(sheet.name, rows)
                pages.append(page)
                ledger.extend(records)
        finally:
            workbook.release_resources()
    if sum(len(p['raw_text']) for p in pages) > MAX_CHARS:
        raise ValueError('Document text exceeds configured character limit; split the document.')
    if not pages or not any(p['raw_text'].strip() for p in pages):
        raise ValueError('No readable text found in document.')
    return {'file_type': kind, 'pages': pages, 'ledger_rows': ledger, 'page_count': len(pages)}
